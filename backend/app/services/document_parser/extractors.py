import io
try:
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    import fitz  # Fallback
import docx
from app.core.exceptions import BadRequestCustomException


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract full structured text from a PDF document using PyMuPDF."""
    text_content = []
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            if doc.page_count == 0:
                raise BadRequestCustomException("The uploaded PDF has no pages.")

            for page_num in range(doc.page_count):
                page = doc.load_page(page_num)
                page_text = page.get_text("text")
                if page_text:
                    text_content.append(page_text.strip())

        full_text = "\n\n".join(text_content).strip()
        if not full_text:
            raise BadRequestCustomException(
                "Could not extract readable text from the PDF. If this is a scanned image, please upload a text-based PDF or DOCX."
            )
        return full_text
    except BadRequestCustomException:
        raise
    except Exception as e:
        raise BadRequestCustomException(f"Failed to process PDF document: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract structured text from a DOCX document using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(" | ".join(row_text))

        full_text = "\n".join(paragraphs).strip()
        if not full_text:
            raise BadRequestCustomException("The uploaded DOCX file contains no readable text.")
        return full_text
    except BadRequestCustomException:
        raise
    except Exception as e:
        raise BadRequestCustomException(f"Failed to process DOCX document: {str(e)}")


def extract_text_from_document(file_bytes: bytes, filename: str) -> str:
    """Detect file type and extract raw text."""
    lower_filename = filename.lower()
    if lower_filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise BadRequestCustomException(
            f"Unsupported file type. Please upload a PDF or DOCX file (got '{filename}')."
        )
